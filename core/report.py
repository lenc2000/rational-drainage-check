"""
수식이 연결된 엑셀 계산서 (한 지점 = 한 파일)

  1_입력 → 2_유출계수 → 3_도달시간 → 4_강우강도 → 5_홍수량 → 6_배수로경사 → 7_단면검토
                                                                         ↑
                                                         8_등류수심 (이분법 계산용)
파란 글씨 = 입력값(바꾸면 전체가 다시 계산됨), 검은 글씨 = 수식, 초록 글씨 = 다른 시트 참조
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

FONT = "Arial"
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="666666")
F_HEAD = Font(name=FONT, size=10, bold=True)
F_IN = Font(name=FONT, size=10, color="0000FF")
F_CALC = Font(name=FONT, size=10, color="000000")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BIG = Font(name=FONT, size=14, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

S1, S2, S3, S4, S5, S6, S7, S8 = ("1_입력", "2_유출계수", "3_도달시간", "4_강우강도", "5_홍수량",
                                  "6_배수로경사", "7_단면검토", "8_등류수심")
LU_ROWS = 12        # 토지이용 입력 행 수
SEG_ROWS = 15       # 배수로 구간 입력 행 수
ITER = 40           # 이분법 반복 횟수


def ref(sheet: str, cell: str) -> str:
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    return f"'{sheet}'!${col}${row}"


def _cell(ws, addr, value, font=F_CALC, fmt=None, fill=None, align=CENTER, border=True):
    c = ws[addr]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = align
    if border:
        c.border = BOX
    return c


def _header(ws, row, labels, col0=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BOX


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, title, sub=None):
    ws["A1"].value, ws["A1"].font = title, F_TITLE
    if sub:
        ws["A2"].value, ws["A2"].font = sub, F_NOTE


# ---------------------------------------------------------------------------
def build_workbook(p: dict) -> bytes:
    wb = Workbook()
    R = {}                                       # 입력 이름 → 절대참조

    # =====================================================================
    # 1_입력
    # =====================================================================
    ws = wb.active
    ws.title = S1
    _title(ws, "합리식 수리계산서 — 입력값",
           "파란 글씨(노란 칸) = 입력값. 값을 바꾸면 모든 시트가 다시 계산됩니다. 검은 글씨 = 수식, 초록 글씨 = 다른 시트 참조.")
    _widths(ws, [30, 34, 12, 70])
    _header(ws, 4, ["항목", "값", "단위", "설명"])
    rows = [
        ("■ 지점·강우", None, None, None, None),
        ("site", "지점명", p["site"], "", "계산서 표지용"),
        ("station", "강우관측소", p["station_label"], "", "WAMIS 강우관측소 (이름, 관할기관, 대유역, 코드)"),
        ("T", "재현기간", p["T"], "년", ""),
        ("a", "강우강도식 계수 a", p["a"], "", "WAMIS 「전국 하천유역 홍수량 및 확률강우량 정보」 강우강도식_IDF_General형 (GEV, 지역빈도해석)"),
        ("b", "강우강도식 계수 b", p["b"], "", "I = a / (t^n + b),  I: mm/hr, t: 분"),
        ("n", "강우강도식 계수 n", p["n"], "", "관측소·재현기간을 바꾸면 a, b, n 도 함께 바꿔야 합니다"),
        ("target", "방재성능목표 1시간 강우량", p["target"], "mm", "해당 시·군 공고값. 0 이면 비교 생략"),
        ("apply", "방재성능목표 보정 반영", "예" if p["apply_target"] else "아니오", "", "예 = 목표가 강우강도식 1시간 강우량보다 클 때 그 비율만큼 강우강도를 키움"),
        ("■ 유역", None, None, None, None),
        ("A", "유역면적", p["A_ha"], "ha", "유출계수는 2_유출계수 시트에서 입력"),
        ("L", "유로연장", p["L"], "m", "최원점 ~ 출구"),
        ("H", "표고차", p["H"], "m", "최원점 표고 − 출구 표고"),
        ("■ 도달시간", None, None, None, None),
        ("method", "도달시간 산정방법", p["tc_method"], "",
         "kraven2 = Kraven(II) 경사별 유속표 / kraven_cont = Kraven 연속식(기존 양식) / rziha / kirpich / velocity = 유속 직접 / direct = tc 직접"),
        ("inlet_mode", "유입시간 방식", "Kerby" if p["inlet_mode"] == "kerby" else "직접", "", "Kerby 또는 직접 (kirpich·direct 방법에서는 쓰지 않음)"),
        ("t1_direct", "유입시간 (직접 입력)", p["t1_direct"], "분", "유입시간 방식 = 직접 일 때"),
        ("k_L", "Kerby 표면류 길이", p["kerby_L"], "m", "t1 = 1.44 × (L·N/√S)^0.467"),
        ("k_N", "Kerby 지체계수 N", p["kerby_N"], "", "포장 0.02, 나지 0.1, 초지 0.3~0.4, 산림 0.6~0.8"),
        ("k_S", "Kerby 사면경사", p["kerby_S"], "m/m", ""),
        ("V_direct", "유하 평균유속 (직접 입력)", p["V_direct"], "m/s", "방법 = velocity 일 때"),
        ("tc_direct", "도달시간 (직접 입력)", p["tc_direct"], "분", "방법 = direct 일 때"),
        ("■ 배수로 단면 (구간별 경사는 6_배수로경사 시트에서 입력)", None, None, None, None),
        ("shape", "단면 형식", "사각" if p["shape"] == "rect" else "원형", "", "사각 또는 원형"),
        ("B", "폭 B (사각)", p["B"], "m", ""),
        ("Hs", "높이 H (사각)", p["Hs"], "m", ""),
        ("D", "관경 D (원형)", p["D"], "m", "예: D800 → 0.8"),
        ("rough", "조도계수 n", p["rough"], "", "콘크리트 측구 0.015, 흄관 0.013, PE관 0.011 등"),
        ("allow", "허용 수심비", p["allow"], "", "계획홍수량(단면 통수량)을 이 수심에서 계산 (실무 80%)"),
        ("safety", "홍수량 여유율", p["safety"], "배", "1.0 이면 여유율 없이 비교"),
        ("vmin", "최소유속", p["vmin"], "m/s", ""),
        ("vmax", "최대유속", p["vmax"], "m/s", ""),
    ]
    r = 5
    for key, lab, val, unit, note in rows:
        if lab is None:
            ws.cell(row=r, column=1, value=key).font = F_SUB
            r += 1
            continue
        _cell(ws, f"A{r}", lab, F_CALC, align=LEFT)
        c = _cell(ws, f"B{r}", val, F_IN, fill=FILL_IN)
        if isinstance(val, float):
            c.number_format = "0.000####"
        _cell(ws, f"C{r}", unit)
        _cell(ws, f"D{r}", note, F_NOTE, align=LEFT)
        R[key] = ref(S1, f"B{r}")
        r += 1

    def _dv(key, options):
        cell = R[key].split("!")[1].replace("$", "")
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(cell)
    _dv("apply", ["예", "아니오"])
    _dv("shape", ["사각", "원형"])
    _dv("inlet_mode", ["Kerby", "직접"])
    _dv("method", ["kraven2", "kraven_cont", "rziha", "kirpich", "velocity", "direct"])

    # =====================================================================
    # 2_유출계수
    # =====================================================================
    ws = wb.create_sheet(S2)
    _title(ws, "2. 유출계수 (토지이용별 면적가중평균)", "토지이용, 면적, 유출계수를 입력하세요. 빈 행은 계산에서 빠집니다.")
    _widths(ws, [20, 16, 12, 16])
    _header(ws, 4, ["토지이용", "면적 (㎡)", "유출계수 C", "C × 면적"])
    lu = p["landuse"]
    r0 = 5
    for i in range(LU_ROWS):
        rr = r0 + i
        name, area, cval = lu[i] if i < len(lu) else (None, None, None)
        _cell(ws, f"A{rr}", name, F_IN, fill=FILL_IN)
        _cell(ws, f"B{rr}", area, F_IN, "#,##0.0", FILL_IN)
        _cell(ws, f"C{rr}", cval, F_IN, "0.000", FILL_IN)
        _cell(ws, f"D{rr}", f'=IF(OR(B{rr}="",C{rr}=""),"",B{rr}*C{rr})', fmt="#,##0.0")
    r1 = r0 + LU_ROWS - 1
    rt = r1 + 1
    _cell(ws, f"A{rt}", "합계", F_HEAD)
    _cell(ws, f"B{rt}", f"=SUM(B{r0}:B{r1})", fmt="#,##0.0")
    _cell(ws, f"C{rt}", "")
    _cell(ws, f"D{rt}", f"=SUM(D{r0}:D{r1})", fmt="#,##0.0")
    _cell(ws, f"A{rt + 2}", "가중평균 유출계수 C", F_HEAD, align=LEFT)
    _cell(ws, f"B{rt + 2}", f"=D{rt}/B{rt}", F_BIG, "0.000", FILL_KEY)
    R["C"] = ref(S2, f"B{rt + 2}")
    _cell(ws, f"A{rt + 3}", "입력 유역면적 (㎡)", align=LEFT)
    _cell(ws, f"B{rt + 3}", f"={R['A']}*10000", F_LINK, "#,##0.0")
    _cell(ws, f"A{rt + 4}", "면적 차이 (토지이용 합계 − 유역면적)", align=LEFT)
    _cell(ws, f"B{rt + 4}", f"=B{rt}-B{rt + 3}", fmt="#,##0.0")
    _cell(ws, f"C{rt + 4}", "0 이 아니면 확인", F_NOTE, border=False)

    # =====================================================================
    # 3_도달시간
    # =====================================================================
    ws = wb.create_sheet(S3)
    _title(ws, "3. 도달시간", "tc = 유입시간 + 유하시간 (Kirpich, 직접 입력은 tc 그대로)")
    _widths(ws, [34, 16, 10, 70])
    _header(ws, 4, ["항목", "값", "단위", "식"])
    m = R["method"]
    lines = [
        ("S", "유역 평균경사 S = H / L", f"={R['H']}/{R['L']}", "m/m", "0.0000", "표고차 / 유로연장"),
        ("t1", "유입시간 t1",
         f'=IF(OR({m}="kirpich",{m}="direct"),0,IF({R["inlet_mode"]}="Kerby",1.44*({R["k_L"]}*{R["k_N"]}/SQRT({R["k_S"]}))^0.467,{R["t1_direct"]}))',
         "분", "0.00", "Kerby: 1.44 × (L·N/√S)^0.467"),
        ("V", "유하 평균유속 V",
         f'=IF({m}="kraven2",IF(B5>1/100,3.5,IF(B5>=1/200,3,2.1)),'
         f'IF({m}="kraven_cont",IF(B5>3/400,MIN(4.5,4.592-0.01194/B5),MAX(1.6,35151.515*B5^2-79.393939*B5+1.6181818)),'
         f'IF({m}="rziha",72*B5^0.6/3.6,IF({m}="velocity",{R["V_direct"]},0))))',
         "m/s", "0.00", "Kraven(II): S>1/100 → 3.5, 1/200~1/100 → 3.0, 그 이하 2.1 / 연속식 / Rziha: 72·S^0.6 km/hr"),
        ("t2", "유하시간 t2 = L / V / 60",
         f'=IF({m}="kirpich",3.976*({R["L"]}/1000)^0.77*B5^-0.385,IF({m}="direct",0,{R["L"]}/B7/60))',
         "분", "0.00", "Kirpich 는 전체 도달시간 = 3.976 × L(km)^0.77 × S^-0.385"),
        ("tc", "도달시간 tc", f'=IF({m}="direct",{R["tc_direct"]},B6+B8)', "분", "0.00", ""),
    ]
    for i, (key, lab, f, unit, fmt, note) in enumerate(lines):
        rr = 5 + i
        _cell(ws, f"A{rr}", lab, align=LEFT)
        _cell(ws, f"B{rr}", f, F_BIG if key == "tc" else F_CALC, fmt, FILL_KEY if key == "tc" else None)
        _cell(ws, f"C{rr}", unit)
        _cell(ws, f"D{rr}", note, F_NOTE, align=LEFT)
        R[key] = ref(S3, f"B{rr}")
    _cell(ws, "A11", "방법", align=LEFT)
    _cell(ws, "B11", f"={m}", F_LINK)

    # =====================================================================
    # 4_강우강도
    # =====================================================================
    ws = wb.create_sheet(S4)
    _title(ws, "4. 강우강도 및 방재성능목표 비교", "I = a / (t^n + b).  방재성능목표는 1시간 평균값이므로 같은 비율로 도달시간 강우강도를 보정")
    _widths(ws, [40, 16, 10, 60])
    _header(ws, 4, ["항목", "값", "단위", "비고"])
    a, b, n = R["a"], R["b"], R["n"]
    lines = [
        ("I_tc", "도달시간 강우강도 I(tc)", f"={a}/({R['tc']}^{n}+{b})", "mm/hr", "0.00", ""),
        ("I60", "강우강도식 1시간 강우량 I(60)×1", f"={a}/(60^{n}+{b})", "mm", "0.00", "1시간 강도(mm/hr) = 1시간 강우량(mm)"),
        ("tgt", "방재성능목표 1시간 강우량", f"={R['target']}", "mm", "0.00", ""),
        ("ratio", "목표 / 강우강도식 1시간", '=IF(B7>0,B7/B6,"-")', "", "0.000", "1 보다 크면 목표가 더 큼"),
        ("I_t", "목표 비율로 보정한 강우강도", '=IF(B7>0,B5*B8,"-")', "mm/hr", "0.00", ""),
        ("I", "설계 강우강도 I", f'=IF(AND({R["apply"]}="예",B7>0,B8>1),B9,B5)', "mm/hr", "0.00", ""),
        ("basis", "설계 강우강도 근거", f'=IF(AND({R["apply"]}="예",B7>0,B8>1),"방재성능목표 보정","확률강우강도식")', "", None, ""),
    ]
    for i, (key, lab, f, unit, fmt, note) in enumerate(lines):
        rr = 5 + i
        _cell(ws, f"A{rr}", lab, align=LEFT)
        _cell(ws, f"B{rr}", f, F_BIG if key == "I" else F_CALC, fmt, FILL_KEY if key == "I" else None)
        _cell(ws, f"C{rr}", unit)
        _cell(ws, f"D{rr}", note, F_NOTE, align=LEFT)
        R[key] = ref(S4, f"B{rr}")
    _cell(ws, "A13", "※ WAMIS 확률강우량은 1시간 이상 자료로 산정되어, 1시간 미만 지속기간은 식을 연장해 쓰는 것입니다.",
          F_NOTE, align=LEFT, border=False)

    # =====================================================================
    # 5_홍수량
    # =====================================================================
    ws = wb.create_sheet(S5)
    _title(ws, "5. 첨두홍수량 (합리식)", "Q = (1/3.6) × C × I × A   [Q: ㎥/s, I: mm/hr, A: ㎢]")
    _widths(ws, [34, 16, 10, 40])
    _header(ws, 4, ["항목", "값", "단위", "비고"])
    lines = [
        ("A_km2", "유역면적 A", f"={R['A']}/100", "㎢", "0.000000", "ha / 100"),
        ("C_", "유출계수 C", f"={R['C']}", "", "0.000", "2_유출계수"),
        ("I_", "설계 강우강도 I", f"={R['I']}", "mm/hr", "0.00", "4_강우강도"),
        ("Q", "첨두홍수량 Q", "=B6*B7*B5/3.6", "㎥/s", "0.0000", "식: 0.2778 × C × I × A"),
    ]
    for i, (key, lab, f, unit, fmt, note) in enumerate(lines):
        rr = 5 + i
        _cell(ws, f"A{rr}", lab, align=LEFT)
        _cell(ws, f"B{rr}", f, F_BIG if key == "Q" else F_LINK if key in ("C_", "I_") else F_CALC,
              fmt, FILL_KEY if key == "Q" else None)
        _cell(ws, f"C{rr}", unit)
        _cell(ws, f"D{rr}", note, F_NOTE, align=LEFT)
        R[key] = ref(S5, f"B{rr}")

    # =====================================================================
    # 6_배수로경사
    # =====================================================================
    ws = wb.create_sheet(S6)
    _title(ws, "6. 배수로 구간별 경사 → 등가경사",
           "구간별 연장과 시점·종점 표고를 입력하세요. 등가경사 S = (ΣL / Σ(L/√S))²  (구간별 유하시간 합이 같아지는 단일 경사)")
    _widths(ws, [12, 12, 14, 14, 12, 12, 14])
    _header(ws, 4, ["구간", "연장 L (m)", "시점 표고 (m)", "종점 표고 (m)", "낙차 (m)", "경사 S", "L / √S"])
    segs = p["segments"]
    r0 = 5
    for i in range(SEG_ROWS):
        rr = r0 + i
        nm, L, z1, z2 = segs[i] if i < len(segs) else (None, None, None, None)
        _cell(ws, f"A{rr}", nm, F_IN, fill=FILL_IN)
        _cell(ws, f"B{rr}", L, F_IN, "0.00", FILL_IN)
        _cell(ws, f"C{rr}", z1, F_IN, "0.000", FILL_IN)
        _cell(ws, f"D{rr}", z2, F_IN, "0.000", FILL_IN)
        _cell(ws, f"E{rr}", f'=IF(B{rr}="","",C{rr}-D{rr})', fmt="0.000")
        _cell(ws, f"F{rr}", f'=IF(B{rr}="","",E{rr}/B{rr})', fmt="0.00000")
        _cell(ws, f"G{rr}", f'=IF(B{rr}="","",IF(F{rr}>0,B{rr}/SQRT(F{rr}),"경사오류"))', fmt="0.00")
    r1 = r0 + SEG_ROWS - 1
    rt = r1 + 1
    _cell(ws, f"A{rt}", "합계", F_HEAD)
    _cell(ws, f"B{rt}", f"=SUM(B{r0}:B{r1})", fmt="0.00")
    _cell(ws, f"E{rt}", f"=SUM(E{r0}:E{r1})", fmt="0.000")
    _cell(ws, f"G{rt}", f"=SUM(G{r0}:G{r1})", fmt="0.00")
    out = [
        ("S_avg", "단순 평균경사 (총낙차/총연장)", f"=E{rt}/B{rt}", "0.00000"),
        ("S_eq", "등가경사 (설계 적용)", f'=IF(COUNTIF(G{r0}:G{r1},"경사오류")>0,"경사오류",(B{rt}/G{rt})^2)', "0.00000"),
        ("S_min", "가장 완만한 구간 경사", f"=MIN(F{r0}:F{r1})", "0.00000"),
        ("S_max", "가장 급한 구간 경사", f"=MAX(F{r0}:F{r1})", "0.00000"),
    ]
    for i, (key, lab, f, fmt) in enumerate(out):
        rr = rt + 2 + i
        ws.merge_cells(f"A{rr}:C{rr}")
        _cell(ws, f"A{rr}", lab, F_HEAD if key == "S_eq" else F_CALC, align=LEFT)
        _cell(ws, f"D{rr}", f, F_BIG if key == "S_eq" else F_CALC, fmt, FILL_KEY if key == "S_eq" else None)
        _cell(ws, f"E{rr}", f'=IF(ISNUMBER(D{rr}),D{rr},"")', fmt="0.00%")
        R[key] = ref(S6, f"D{rr}")
    ws[f"F{rt + 3}"].value = "← 7_단면검토의 설계 경사"
    ws[f"F{rt + 3}"].font = F_NOTE

    # =====================================================================
    # 8_등류수심 (이분법) — 7_단면검토에서 참조
    # =====================================================================
    ws8 = wb.create_sheet(S8)
    _title(ws8, "8. 등류수심 계산 (이분법, 계산용 시트)",
           f"Manning 식을 만족하는 수심을 {ITER}번 반으로 좁혀 구합니다. 직접 고칠 필요 없습니다.")
    shape, B_, Hs_, D_, nr = R["shape"], R["B"], R["Hs"], R["D"], R["rough"]
    Qd = ref(S7, "C9")                       # 설계 비교유량 (7_단면검토)

    def q_of(ycell, scell):
        th = f"2*ACOS(1-2*{ycell}/{D_})"
        a_c = f"({D_}^2/8*({th}-SIN({th})))"
        p_c = f"({D_}*{th}/2)"
        a_r = f"({B_}*{ycell})"
        p_r = f"({B_}+2*{ycell})"
        return (f'=IF({shape}="사각",{a_r}*({a_r}/{p_r})^(2/3),{a_c}*({a_c}/{p_c})^(2/3))'
                f"*SQRT({scell})/{nr}")

    blocks = [("설계 경사 (등가경사)", R["S_eq"], 1), ("가장 급한 구간 경사", R["S_max"], 7)]
    for title, sref, c0 in blocks:
        cl = [get_column_letter(c0 + k) for k in range(5)]   # 반복, 하한, 상한, 중간, Q(중간)
        ws8.cell(row=4, column=c0, value=title).font = F_SUB
        ws8.cell(row=5, column=c0, value="경사").font = F_HEAD
        ws8.cell(row=5, column=c0 + 1, value=f'=IF(ISNUMBER({sref}),{sref},0.0001)').font = F_LINK
        ws8.cell(row=6, column=c0, value="최대수심").font = F_HEAD
        ws8.cell(row=6, column=c0 + 1, value=f'=IF({shape}="사각",{Hs_},0.938*{D_})')
        ws8.cell(row=7, column=c0, value="최대수심 유량").font = F_HEAD
        ws8.cell(row=7, column=c0 + 1, value=q_of(f"{cl[1]}6", f"{cl[1]}5")[0:])
        _header(ws8, 9, ["반복", "하한", "상한", "중간", "Q(중간)"], col0=c0)
        for k in range(ITER):
            rr = 10 + k
            ws8[f"{cl[0]}{rr}"] = k + 1
            if k == 0:
                ws8[f"{cl[1]}{rr}"] = 0
                ws8[f"{cl[2]}{rr}"] = f"={cl[1]}6"
            else:
                ws8[f"{cl[1]}{rr}"] = f"=IF({cl[4]}{rr - 1}<{Qd},{cl[3]}{rr - 1},{cl[1]}{rr - 1})"
                ws8[f"{cl[2]}{rr}"] = f"=IF({cl[4]}{rr - 1}<{Qd},{cl[2]}{rr - 1},{cl[3]}{rr - 1})"
            ws8[f"{cl[3]}{rr}"] = f"=({cl[1]}{rr}+{cl[2]}{rr})/2"
            ws8[f"{cl[4]}{rr}"] = q_of(f"{cl[3]}{rr}", f"{cl[1]}$5")
            for c in cl:
                ws8[f"{c}{rr}"].font = F_CALC
                ws8[f"{c}{rr}"].number_format = "0.000000"
        last = 10 + ITER - 1
        R[f"y_{c0}"] = f"'{S8}'!${cl[3]}${last}"
        R[f"qmax_{c0}"] = f"'{S8}'!${cl[1]}$7"
    for col in range(1, 13):
        ws8.column_dimensions[get_column_letter(col)].width = 13

    # =====================================================================
    # 7_단면검토
    # =====================================================================
    ws = wb.create_sheet(S7, index=6)
    _title(ws, "7. 배수로 단면 검토 — 홍수량 vs 계획홍수량(단면 통수량)",
           "계획홍수량 = 허용 수심(기본 80%)에서 단면이 흘릴 수 있는 유량. 홍수량 ≤ 계획홍수량 이면 통수 가능")
    _widths(ws, [34, 18, 18, 12, 50])
    _header(ws, 4, ["항목", "", "값", "단위", "비고"])
    yA = "C13"
    th_a = f"2*ACOS(1-2*{yA}/{D_})"
    rows7 = [
        (5, "단면", f'=IF({shape}="사각","사각 B "&TEXT({B_},"0.00")&" × H "&TEXT({Hs_},"0.00")&" m","원형 D "&TEXT({D_}*1000,"0")&" mm")', "", None, ""),
        (6, "조도계수 n", f"={nr}", "", "0.000", ""),
        (7, "설계 경사 S (등가경사)", f"={R['S_eq']}", "m/m", "0.00000", "6_배수로경사"),
        (8, "홍수량 Q", f"={R['Q']}", "㎥/s", "0.0000", "5_홍수량"),
        (9, "비교 홍수량 Q × 여유율", f"=C8*{R['safety']}", "㎥/s", "0.0000", ""),
        (11, "■ 계획홍수량 (허용 수심에서의 단면 통수량)", None, None, None, None),
        (12, "단면 높이 (사각 H / 원형 D)", f'=IF({shape}="사각",{Hs_},{D_})', "m", "0.000", ""),
        (13, "허용 수심 y", f"={R['allow']}*C12", "m", "0.000", "높이 × 허용 수심비"),
        (14, "통수 단면적 A", f'=IF({shape}="사각",{B_}*{yA},{D_}^2/8*({th_a}-SIN({th_a})))', "㎡", "0.0000", ""),
        (15, "윤변 P", f'=IF({shape}="사각",{B_}+2*{yA},{D_}*({th_a})/2)', "m", "0.0000", ""),
        (16, "동수반경 R = A / P", "=C14/C15", "m", "0.0000", ""),
        (17, "평균유속 V (허용 수심)", f"=C16^(2/3)*SQRT(C7)/{nr}", "m/s", "0.000", "Manning: V = (1/n)·R^(2/3)·S^(1/2)"),
        (18, "계획홍수량 Qa = A × V", "=C14*C17", "㎥/s", "0.0000", ""),
        (20, "■ 판정", None, None, None, None),
        (21, "차이 (계획홍수량 − 홍수량)", "=C18-C9", "㎥/s", "0.0000", ""),
        (22, "여유", "=C18/C9-1", "", "0.0%", ""),
        (23, "통수 판정", '=IF(C18>=C9,"통수 가능 (OK)","통수 불가 (NG)")', "", None, ""),
        (25, "■ 유속 확인 (별도)", None, None, None, None),
        (26, "등류수심 (설계 경사)", f'=IF({R["qmax_1"]}<C9,"단면 초과",{R["y_1"]})', "m", "0.000", "8_등류수심 이분법"),
        (27, "수심비", '=IF(ISNUMBER(C26),C26/C12,"-")', "", "0.0%", ""),
        (28, "유속 V (설계 경사)",
         f'=IF(ISNUMBER(C26),C9/IF({shape}="사각",{B_}*C26,{D_}^2/8*(2*ACOS(1-2*C26/{D_})-SIN(2*ACOS(1-2*C26/{D_})))),"-")',
         "m/s", "0.00", ""),
        (29, "유속 판정",
         f'=IF(ISNUMBER(C28),IF(C28<{R["vmin"]},"최소유속 미달 (퇴적 우려)",IF(C28>{R["vmax"]},"최대유속 초과 (세굴·마모 우려)","허용범위 이내 (OK)")),"-")',
         "", None, ""),
        (31, "■ 구간별 참고", None, None, None, None),
        (32, "가장 완만한 구간 계획홍수량",
         f"=C14*C16^(2/3)*SQRT(MAX({R['S_min']},0.000001))/{nr}", "㎥/s", "0.0000", "가장 완만한 구간에서도 홍수량을 흘리는지"),
        (33, "가장 완만한 구간 판정", '=IF(C32>=C9,"통수 가능","통수 불가")', "", None, ""),
        (34, "가장 급한 구간 유속",
         f'=IF({R["qmax_7"]}<C9,"단면 초과",C9/IF({shape}="사각",{B_}*{R["y_7"]},{D_}^2/8*(2*ACOS(1-2*{R["y_7"]}/{D_})-SIN(2*ACOS(1-2*{R["y_7"]}/{D_})))))',
         "m/s", "0.00", "가장 급한 구간에서 유속이 최대가 됨"),
        (35, "가장 급한 구간 유속 판정",
         f'=IF(ISNUMBER(C34),IF(C34>{R["vmax"]},"최대유속 초과",IF(C34<{R["vmin"]},"최소유속 미달","허용범위 이내")),"-")',
         "", None, ""),
    ]
    for rr, lab, f, unit, fmt, note in rows7:
        if f is None:
            ws.cell(row=rr, column=1, value=lab).font = F_SUB
            continue
        ws.merge_cells(f"A{rr}:B{rr}")
        _cell(ws, f"A{rr}", lab, align=LEFT)
        key = rr in (18, 23, 8)
        _cell(ws, f"C{rr}", f, F_BIG if key else (F_LINK if rr in (6, 7, 8) else F_CALC), fmt,
              FILL_KEY if key else None)
        _cell(ws, f"D{rr}", unit)
        _cell(ws, f"E{rr}", note, F_NOTE, align=LEFT)
    ws["C18"].comment = Comment("계획홍수량: 허용 수심(1_입력의 허용 수심비)에서 이 단면이 흘릴 수 있는 유량", "계산서")

    # 요약을 1_입력 아래에도 표시
    ws1 = wb[S1]
    base = ws1.max_row + 2
    ws1.cell(row=base, column=1, value="■ 결과 요약").font = F_SUB
    summ = [("도달시간 tc (분)", f"={R['tc']}", "0.00"), ("설계 강우강도 I (mm/hr)", f"={R['I']}", "0.00"),
            ("홍수량 Q (㎥/s)", f"='{S7}'!C8", "0.0000"), ("계획홍수량 Qa (㎥/s)", f"='{S7}'!C18", "0.0000"),
            ("차이 (㎥/s)", f"='{S7}'!C21", "0.0000"), ("통수 판정", f"='{S7}'!C23", None),
            ("유속 (m/s)", f"='{S7}'!C28", "0.00"), ("유속 판정", f"='{S7}'!C29", None)]
    for i, (lab, f, fmt) in enumerate(summ):
        rr = base + 1 + i
        _cell(ws1, f"A{rr}", lab, align=LEFT)
        _cell(ws1, f"B{rr}", f, F_LINK, fmt, FILL_KEY if "판정" in lab else None)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
